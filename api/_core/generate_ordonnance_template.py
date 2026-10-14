import argparse
import json
import textwrap
from docx import Document
from docx.shared import Pt


def clear_paragraph(paragraph):
    element = paragraph._element
    for child in list(element):
        if child.tag.endswith('}r') or child.tag.endswith('}hyperlink'):
            element.remove(child)


def add_run(paragraph, text, *, bold=False, italic=False, size=12):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    return run


def set_prescription_paragraph(paragraph, text, *, bold=False):
    clear_paragraph(paragraph)
    if not text:
        return
    add_run(paragraph, text, bold=bold, size=12)


def wrap_lines(lines, width=78):
    wrapped = []
    for raw in lines:
        value = str(raw or "").strip()
        if not value:
            continue
        wrapped.extend(textwrap.wrap(value, width=width) or [""])
    return wrapped


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--payload", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    with open(args.payload, "r", encoding="utf-8-sig") as fh:
        payload = json.load(fh)

    doc = Document(args.template)
    paragraphs = doc.paragraphs

    # P1 doctor block
    para = paragraphs[1]
    clear_paragraph(para)
    add_run(para, payload.get("doctor_name", "Dr. Medecin"), bold=True, size=12)
    add_run(para, "\t\t\t\t\t\t", size=12)
    add_run(para, payload.get("doctor_address_line1", ""), size=12)
    add_run(para, "\n", size=12)
    add_run(para, payload.get("doctor_specialite", "Medecin"), size=12)
    add_run(para, "\t\t\t\t\t\t", size=12)
    add_run(para, payload.get("doctor_address_line2", ""), size=12)

    # P3 secondary line and phone
    para = paragraphs[3]
    clear_paragraph(para)
    add_run(para, payload.get("doctor_secondary_line", ""), italic=True, size=12)
    add_run(para, "\t\t\t", italic=True, size=12)
    add_run(para, payload.get("doctor_phone", ""), italic=True, size=12)

    # P6 patient block
    para = paragraphs[6]
    clear_paragraph(para)
    patient_name = payload.get("patient_name", "Patient")
    meta_lines = payload.get("patient_meta_lines", []) or []
    add_run(para, patient_name, bold=True, size=12)
    for line in meta_lines:
        add_run(para, "\n", size=12)
        add_run(para, line, size=12)

    # P7 issue place/date
    para = paragraphs[7]
    clear_paragraph(para)
    add_run(para, payload.get("issue_place_date", ""), size=12)

    # Prescription area P10..P19
    prescription_title = str(payload.get("prescription_title", "") or "").strip()
    prescription_lines = wrap_lines(payload.get("prescription_lines", []) or [])
    slots = list(range(10, 20))

    set_prescription_paragraph(paragraphs[10], prescription_title, bold=bool(prescription_title))
    regular_slots = slots[1:] if prescription_title else slots
    regular_index = 0
    for slot in regular_slots:
        para = paragraphs[slot]
        if regular_index < len(prescription_lines):
            text = prescription_lines[regular_index]
            regular_index += 1
            while regular_index < len(prescription_lines) and slot == regular_slots[-1]:
                text += "\n" + prescription_lines[regular_index]
                regular_index += 1
            clear_paragraph(para)
            for idx, subline in enumerate(text.split("\n")):
                if idx > 0:
                    add_run(para, "\n", size=12)
                add_run(para, subline, size=12)
        else:
            clear_paragraph(para)

    doc.save(args.output)


if __name__ == "__main__":
    main()
